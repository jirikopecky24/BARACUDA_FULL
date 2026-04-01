from __future__ import annotations

"""XLSX export for run trajectories.

Goals:
- Human-friendly sheet with freeze header + filter.
- Metadata sheet with structured metadata + raw header lines from trajectory.csv.
- Never be the reason a run fails (caller should catch exceptions).

This exporter is tolerant to extra columns (e.g. OT-3.1 adds lost/drift columns).
"""

import math
from pathlib import Path
from typing import Any

from barakuda.core.ot_report import build_ot_item_summary, build_ot_summary_rows
from barakuda.core.trajectory_csv_io import read_trajectory_csv


def export_trajectory_xlsx(
    run_dir: Path,
    trajectory_csv_path: Path,
    extra_metadata: dict[str, Any] | None = None,
    output_name: str = "trajectory.xlsx",
) -> Path:
    """Create run_dir/output_name with sheets Trajectory + Metadata.

    Requires openpyxl.
    """

    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except Exception as e:
        raise RuntimeError("openpyxl is required for XLSX export. Install it via pip/conda.") from e

    run_dir = Path(run_dir)
    trajectory_csv_path = Path(trajectory_csv_path)
    out_path = run_dir / output_name

    table = read_trajectory_csv(trajectory_csv_path)
    header = list(table.header)

    wb = Workbook()

    # --- Sheet: Trajectory ---
    ws = wb.active
    ws.title = "Trajectory"

    ws.append(header)
    for r in table.rows:
        ws.append([r.get(col, "") for col in header])

    # Usability
    ws.freeze_panes = "A2"
    if header:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}1"

    # Column widths (simple heuristic)
    for i, col in enumerate(header, start=1):
        # Use column name length and sample a few rows
        max_len = len(str(col))
        for rr in table.rows[:200]:
            v = rr.get(col, "")
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        # Clamp to readable range
        ws.column_dimensions[get_column_letter(i)].width = float(max(10, min(28, max_len + 2)))

    # --- Sheet: Metadata ---
    ws2 = wb.create_sheet("Metadata")
    ws2.append(["Key", "Value"])

    if extra_metadata:
        for k in sorted(extra_metadata.keys()):
            v = extra_metadata[k]
            ws2.append([str(k), "" if v is None else str(v)])

    if table.meta_lines:
        ws2.append([])
        ws2.append(["trajectory.csv header lines", ""])
        for line in table.meta_lines:
            ws2.append([line, ""])

    ws2.freeze_panes = "A2"
    ws2.column_dimensions["A"].width = 34
    ws2.column_dimensions["B"].width = 90

    wb.save(out_path)
    return out_path


def export_ot_results_xlsx(
    output_dir: Path,
    base_name: str,
    trajectory_csv_path: Path,
    msd_csv_path: Path | None = None,
    psd_x_csv_path: Path | None = None,
    psd_y_csv_path: Path | None = None,
) -> Path:
    """Create <base_name>_results.xlsx as a human-friendly workbook bundle."""
    import csv
    import json
    import re

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment
        from openpyxl.chart import BarChart, Reference, ScatterChart, Series
        from openpyxl.utils import get_column_letter
    except Exception as e:
        raise RuntimeError("openpyxl is required for XLSX export.") from e

    output_dir = Path(output_dir)
    run_dir = output_dir.parent
    trajectory_csv_path = Path(trajectory_csv_path)
    out_path = output_dir / f"{base_name}_results.xlsx"

    wb = Workbook()
    first_sheet_used = False
    used_titles: set[str] = set()
    sheet_by_key: dict[str, Any] = {}
    added_csv_paths: set[Path] = set()
    numeric_pattern = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")
    invalid_sheet_chars = re.compile(r"[\[\]:*?/\\]")

    def _flatten(prefix: str, obj: Any, out: list[tuple[str, str]]) -> None:
        if isinstance(obj, dict):
            for k in sorted(obj.keys(), key=lambda x: str(x)):
                _flatten(f"{prefix}{k}.", obj[k], out)
        elif isinstance(obj, list):
            out.append((prefix[:-1] if prefix.endswith(".") else prefix, json.dumps(obj, ensure_ascii=False)))
        else:
            key = prefix[:-1] if prefix.endswith(".") else prefix
            out.append((key, "" if obj is None else str(obj)))

    def _coerce_cell(value: str) -> Any:
        text = str(value).strip()
        if text == "":
            return ""
        lower = text.lower()
        if lower in {"nan", "+nan", "-nan", "inf", "+inf", "-inf"}:
            return text
        if re.fullmatch(r"[+-]?\d+", text):
            try:
                return int(text)
            except Exception:
                return text
        if numeric_pattern.fullmatch(text):
            try:
                num = float(text)
                if math.isfinite(num):
                    return num
            except Exception:
                return text
        return text

    def _safe_sheet_title(candidate: str) -> str:
        candidate = invalid_sheet_chars.sub("_", str(candidate)).strip()
        candidate = candidate or "Sheet"
        candidate = candidate[:31]
        title = candidate
        suffix = 2
        while title in used_titles:
            suffix_text = f"_{suffix}"
            title = f"{candidate[: max(1, 31 - len(suffix_text))]}{suffix_text}"
            suffix += 1
        used_titles.add(title)
        return title

    def _new_sheet(title: str):
        nonlocal first_sheet_used
        sheet_title = _safe_sheet_title(title)
        if not first_sheet_used:
            ws = wb.active
            first_sheet_used = True
            ws.title = sheet_title
            return ws
        return wb.create_sheet(sheet_title)

    def _format_sheet(ws) -> None:
        if ws.max_row >= 1:
            ws.freeze_panes = "A2"
            if ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
        for col_idx in range(1, ws.max_column + 1):
            max_len = 10
            for cell_tuple in ws.iter_rows(min_col=col_idx, max_col=col_idx, max_row=min(200, ws.max_row)):
                value = cell_tuple[0].value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = float(min(40, max_len + 2))

    def _read_csv_rows(csv_path: Path, *, skip_comments: bool = False) -> list[list[Any]]:
        rows: list[list[Any]] = []
        with csv_path.open("r", encoding="utf-8", newline="") as f:
            reader = csv.reader(f)
            for row in reader:
                if skip_comments and row and str(row[0]).startswith("#"):
                    continue
                rows.append([_coerce_cell(v) for v in row])
        return rows

    def _add_csv_sheet(title: str, csv_path: Path | None, *, skip_comments: bool = False):
        if csv_path is None:
            return None
        csv_path = Path(csv_path)
        if not csv_path.exists():
            return None
        rows = _read_csv_rows(csv_path, skip_comments=skip_comments)
        if not rows:
            return None
        ws = _new_sheet(title)
        for row in rows:
            ws.append(row)
        _format_sheet(ws)
        added_csv_paths.add(csv_path.resolve())
        return ws

    def _add_pairs_sheet(title: str, pairs: list[tuple[str, Any]]):
        ws = _new_sheet(title)
        ws.append(["Key", "Value"])
        for key, value in pairs:
            ws.append([str(key), "" if value is None else str(value)])
        _format_sheet(ws)
        return ws

    def _header_map(ws) -> dict[str, int]:
        out: dict[str, int] = {}
        for idx in range(1, ws.max_column + 1):
            value = ws.cell(row=1, column=idx).value
            if value is None:
                continue
            out[str(value).strip().lower()] = idx
        return out

    def _sheet_name_for_path(csv_path: Path) -> str:
        rel_path = csv_path.relative_to(run_dir) if csv_path.is_relative_to(run_dir) else csv_path.name
        if isinstance(rel_path, Path):
            parts = list(rel_path.parts)
        else:
            parts = [str(rel_path)]
        parent = parts[-2] if len(parts) > 1 else "root"
        stem = csv_path.stem
        if stem.startswith(f"{base_name}_"):
            stem = stem[len(base_name) + 1 :]
        candidate = f"{parent}_{stem}".strip("_")
        return candidate.upper()

    def _load_json(path: Path | None) -> dict[str, Any] | None:
        if path is None or not path.exists():
            return None
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return None

    def _first_existing(*paths: Path | None) -> Path | None:
        for path in paths:
            if path is not None and Path(path).exists():
                return Path(path)
        return None

    def _discover_csvs() -> list[Path]:
        matches: list[Path] = []
        seen: set[Path] = set()
        search_roots = [run_dir / "csv", run_dir / "tracking", run_dir / "physics", run_dir / "audit", output_dir, run_dir]
        for root in search_roots:
            if not root.exists() or not root.is_dir():
                continue
            for csv_path in sorted(root.glob("*.csv"), key=lambda p: p.name.lower()):
                name = csv_path.name
                if not (name == "results.csv" or name.startswith(f"{base_name}_")):
                    continue
                resolved = csv_path.resolve()
                if resolved in seen:
                    continue
                seen.add(resolved)
                matches.append(csv_path)
        return matches

    def _add_formula_column(ws, header: str, build_formula) -> int | None:
        if ws.max_row < 2:
            return None
        col_idx = ws.max_column + 1
        ws.cell(row=1, column=col_idx, value=header)
        for row_idx in range(2, ws.max_row + 1):
            formula = build_formula(row_idx)
            if formula:
                ws.cell(row=row_idx, column=col_idx, value=formula)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(18.0, float(len(header) + 2))
        return col_idx

    def _add_scatter_chart(ws, x_col: int, y_cols: list[int], title: str, x_title: str, y_title: str, anchor: str, *, log_x: bool = False, log_y: bool = False) -> None:
        if ws.max_row < 2 or not y_cols:
            return
        chart = ScatterChart()
        chart.title = title
        chart.style = 13
        chart.width = 18
        chart.height = 9
        chart.x_axis.title = x_title
        chart.y_axis.title = y_title
        if log_x:
            chart.x_axis.scaling.logBase = 10
        if log_y:
            chart.y_axis.scaling.logBase = 10
        xvalues = Reference(ws, min_col=x_col, min_row=2, max_row=ws.max_row)
        for y_col in y_cols:
            values = Reference(ws, min_col=y_col, min_row=1, max_row=ws.max_row)
            series = Series(values, xvalues, title_from_data=True)
            series.marker.symbol = "circle"
            series.graphicalProperties.line.width = 18000
            chart.series.append(series)
        ws.add_chart(chart, anchor)

    def _add_bar_chart(ws, category_col: int, data_cols: list[int], title: str, y_title: str, anchor: str) -> None:
        if ws.max_row < 2 or not data_cols:
            return
        chart = BarChart()
        chart.title = title
        chart.style = 10
        chart.width = 18
        chart.height = 9
        chart.y_axis.title = y_title
        categories = Reference(ws, min_col=category_col, min_row=2, max_row=ws.max_row)
        for data_col in data_cols:
            data = Reference(ws, min_col=data_col, min_row=1, max_row=ws.max_row)
            chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        ws.add_chart(chart, anchor)

    trajectory_name = trajectory_csv_path.name
    derived_csv = _first_existing(
        output_dir / f"{base_name}_derived.csv",
        run_dir / "csv" / f"{base_name}_derived.csv",
        run_dir / "physics" / f"{base_name}_derived.csv",
        run_dir / "tracking" / f"{base_name}_derived.csv",
    )
    calibration_csv = _first_existing(
        output_dir / f"{base_name}_calibration.csv",
        run_dir / "csv" / f"{base_name}_calibration.csv",
        run_dir / "physics" / f"{base_name}_calibration.csv",
        run_dir / "tracking" / f"{base_name}_calibration.csv",
    )
    compare_csv = _first_existing(
        output_dir / f"{base_name}_compare.csv",
        run_dir / "csv" / f"{base_name}_compare.csv",
        run_dir / "physics" / f"{base_name}_compare.csv",
        run_dir / "tracking" / f"{base_name}_compare.csv",
    )
    hist_csvs = {
        tag: _first_existing(
            output_dir / f"{base_name}_hist_{tag}.csv",
            run_dir / "csv" / f"{base_name}_hist_{tag}.csv",
            run_dir / "physics" / f"{base_name}_hist_{tag}.csv",
            run_dir / "tracking" / f"{base_name}_hist_{tag}.csv",
        )
        for tag in ("x", "y", "r")
    }

    summary_payload = build_ot_item_summary(
        run_dir=run_dir,
        base_name=base_name,
        item_id=run_dir.parent.name if run_dir.parent != run_dir else base_name,
        source_input_path=None,
        status="success",
    )
    summary_payload.setdefault("artifacts", {})["xlsx"] = str(out_path)

    ws_summary = _new_sheet("Summary")
    ws_summary.append(["Group", "Metric", "Value", "Unit", "Notes"])
    for group, metric, value, unit, notes in build_ot_summary_rows(summary_payload):
        ws_summary.append([group, metric, value, unit, notes])
    for row in ws_summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _format_sheet(ws_summary)
    sheet_by_key["summary"] = ws_summary

    core_specs = [
        ("trajectory", "Trajectory", _first_existing(trajectory_csv_path, run_dir / "csv" / trajectory_name, run_dir / "tracking" / trajectory_name), True),
        ("msd", "MSD", _first_existing(msd_csv_path, run_dir / "csv" / f"{base_name}_msd.csv", run_dir / "physics" / f"{base_name}_msd.csv"), False),
        ("psd_x", "PSD_X", _first_existing(psd_x_csv_path, run_dir / "csv" / f"{base_name}_psd_x.csv", run_dir / "physics" / f"{base_name}_psd_x.csv"), False),
        ("psd_y", "PSD_Y", _first_existing(psd_y_csv_path, run_dir / "csv" / f"{base_name}_psd_y.csv", run_dir / "physics" / f"{base_name}_psd_y.csv"), False),
        ("derived", "Derived_Physics", derived_csv, False),
        ("calibration", "CALIBRATION", calibration_csv, False),
        ("hist_x", "HIST_X", hist_csvs["x"], False),
        ("hist_y", "HIST_Y", hist_csvs["y"], False),
        ("hist_r", "HIST_R", hist_csvs["r"], False),
        ("compare", "COMPARE", compare_csv, False),
    ]
    for key, title, csv_path, skip_comments in core_specs:
        ws = _add_csv_sheet(title, csv_path, skip_comments=skip_comments)
        if ws is not None:
            sheet_by_key[key] = ws

    for csv_path in _discover_csvs():
        if csv_path.resolve() in added_csv_paths:
            continue
        _add_csv_sheet(_sheet_name_for_path(csv_path), csv_path, skip_comments=False)

    run_json = _first_existing(run_dir / "audit" / "run.json", run_dir / "run.json", output_dir / "run.json")
    post_json = _first_existing(run_dir / "audit" / f"{base_name}_postprocess.json", output_dir / f"{base_name}_postprocess.json")
    psd_fit_json = _first_existing(run_dir / "audit" / f"{base_name}_psd_fit.json", run_dir / "csv" / f"{base_name}_psd_fit.json", run_dir / "tracking" / f"{base_name}_psd_fit.json", output_dir / f"{base_name}_psd_fit.json")
    calibration_json = _first_existing(run_dir / "audit" / f"{base_name}_calibration.json", run_dir / "csv" / f"{base_name}_calibration.json", run_dir / "tracking" / f"{base_name}_calibration.json", output_dir / f"{base_name}_calibration.json")
    drag_summary_json = _first_existing(
        run_dir / "audit" / f"{base_name}_drag_summary.json",
        run_dir / "physics" / f"{base_name}_drag_summary.json",
        output_dir / f"{base_name}_drag_summary.json",
    )
    drag_json = _first_existing(
        run_dir / "audit" / f"{base_name}_drag.json",
        run_dir / "physics" / f"{base_name}_drag.json",
        output_dir / f"{base_name}_drag.json",
    )
    run_protocol_json = _first_existing(
        run_dir / "audit" / "run_protocol.json",
        run_dir / "run_protocol.json",
        output_dir / "run_protocol.json",
    )

    meta_pairs: list[tuple[str, str]] = []
    if run_json is not None:
        payload = _load_json(run_json)
        if payload is not None:
            _flatten("run.", payload, meta_pairs)
        else:
            meta_pairs.append(("run_json_error", "failed to parse run.json"))
    if post_json is not None:
        payload = _load_json(post_json)
        if payload is not None:
            _flatten("postprocess.", payload, meta_pairs)
        else:
            meta_pairs.append(("postprocess_json_error", "failed to parse *_postprocess.json"))
    try:
        table = read_trajectory_csv(trajectory_csv_path)
        for idx, line in enumerate(table.meta_lines):
            meta_pairs.append((f"trajectory_meta[{idx}]", line))
    except Exception:
        meta_pairs.append(("trajectory_meta_error", "failed to read trajectory.csv meta lines"))
    meta_pairs.append(("workbook.csv_sheet_count", str(len(added_csv_paths))))
    ws_meta = _add_pairs_sheet("Metadata", meta_pairs)
    sheet_by_key["metadata"] = ws_meta

    if "calibration" not in sheet_by_key:
        cal_pairs: list[tuple[str, str]] = []
        payload = _load_json(calibration_json)
        if payload is not None:
            _flatten("calibration.", payload, cal_pairs)
        else:
            cal_pairs = [("status", "MISSING"), ("reason", "No calibration artifacts found")]
        sheet_by_key["calibration"] = _add_pairs_sheet("CALIBRATION", cal_pairs)

    drag_source_kind = "none"
    drag_source_path: Path | None = None
    protocol_payload = _load_json(run_protocol_json)
    protocol_provenance = (protocol_payload or {}).get("provenance") or {}
    drag_payload = _load_json(drag_summary_json)
    if drag_payload is not None:
        drag_source_kind = "drag_summary"
        drag_source_path = drag_summary_json
    else:
        drag_payload = _load_json(drag_json)
        if drag_payload is not None:
            drag_source_kind = "legacy_drag_json"
            drag_source_path = drag_json
    if drag_payload is not None:
        drag_pairs: list[tuple[str, str]] = []
        drag_pairs.append(("report_source_kind", drag_source_kind))
        drag_pairs.append(("report_source_path", str(drag_source_path) if drag_source_path is not None else ""))
        if protocol_provenance:
            for key in (
                "current_drag_input_path",
                "current_drag_output_root",
                "current_drag_report_path",
                "brownian_baseline_folder",
                "selected_calibration_path",
                "selected_trajectory_path",
                "selected_timestamps_path",
                "selected_stage_meta_path",
                "selected_stage_trace_path",
            ):
                drag_pairs.append((f"provenance.{key}", protocol_provenance.get(key, "")))
        if drag_source_kind == "legacy_drag_json":
            drag_pairs.append(("warning", "legacy drag source used"))
        _flatten("drag.", drag_payload, drag_pairs)
        sheet_by_key["drag"] = _add_pairs_sheet("DRAG", drag_pairs)

    fit_sheet = _new_sheet("Fits")
    fit_sheet.append(["Source", "Key", "Value", "Equation_or_Note"])
    fit_param_rows: dict[str, dict[str, int]] = {}
    psd_fit_payload = _load_json(psd_fit_json)
    if psd_fit_payload is not None:
        for axis_key, label in (("fit_x", "X"), ("fit_y", "Y")):
            payload = psd_fit_payload.get(axis_key, {})
            if not isinstance(payload, dict):
                continue
            fit_sheet.append(["PSD_FIT", f"{label}.A", payload.get("A", ""), "PSD(f) = A / (fc_hz^2 + f^2) + B"])
            row_a = fit_sheet.max_row
            fit_sheet.append(["PSD_FIT", f"{label}.fc_hz", payload.get("fc_hz", ""), "cutoff frequency"])
            row_fc = fit_sheet.max_row
            fit_sheet.append(["PSD_FIT", f"{label}.B", payload.get("B", ""), "background floor"])
            row_b = fit_sheet.max_row
            fit_sheet.append(["PSD_FIT", f"{label}.formula", "", "PSD(f) = A / (fc_hz^2 + f^2) + B"])
            fit_param_rows[label.lower()] = {"A": row_a, "fc_hz": row_fc, "B": row_b}
        for key, value in sorted((psd_fit_payload.get("psd_params") or {}).items(), key=lambda item: str(item[0])):
            fit_sheet.append(["PSD_PARAMS", str(key), value, "welch PSD configuration"])
    else:
        fit_sheet.append(["PSD_FIT", "status", "missing", "No *_psd_fit.json found"])

    calibration_payload = _load_json(calibration_json)
    if calibration_payload is not None:
        fit_sheet.append([])
        fit_sheet.append(["CALIBRATION", "relation", "", "kappa_pn_per_um = kappa_n_per_m * 1E6"])
        kappa_payload = calibration_payload.get("kappa", {}) if isinstance(calibration_payload, dict) else {}
        if isinstance(kappa_payload, dict):
            for key in ("kappa_x_n_per_m", "kappa_x_pn_per_um", "kappa_y_n_per_m", "kappa_y_pn_per_um", "kappa_iso_ratio"):
                if key in kappa_payload:
                    fit_sheet.append(["CALIBRATION", key, kappa_payload[key], "from calibration.json"])
        diagnostics_payload = calibration_payload.get("diagnostics", {}) if isinstance(calibration_payload, dict) else {}
        if isinstance(diagnostics_payload, dict):
            for key in ("fc_x_hz", "fc_y_hz", "var_x_um2", "var_y_um2"):
                if key in diagnostics_payload:
                    fit_sheet.append(["CALIBRATION", key, diagnostics_payload[key], "calibration diagnostics"])
    if drag_payload is not None:
        fit_sheet.append([])
        if drag_source_kind == "drag_summary":
            fit_sheet.append(["DRAG", "drag_force_n", drag_payload.get("drag_force_n", ""), "drag_force_n = kappa_n_per_m * abs_offset_m"])
            fit_sheet.append(["DRAG", "kappa_pn_per_um", drag_payload.get("kappa_pn_per_um", ""), "baseline calibration stiffness"])
            fit_sheet.append(["DRAG", "abs_offset_um", drag_payload.get("abs_offset_um", ""), "abs_offset_um = |steady_um - baseline_um|"])
            fit_sheet.append(["DRAG", "actual_speed_um_s", drag_payload.get("actual_speed_um_s", ""), "speed inferred from stage timing"])
            fit_sheet.append(["DRAG", "eta_pa_s", drag_payload.get("eta_pa_s", ""), "viscosity inferred from drag force and speed"])
        else:
            fit_sheet.append(["DRAG", "offset_um", drag_payload.get("means_um", {}).get("offset_um", ""), "offset_um = steady_mean_um - baseline_mean_um"])
            fit_sheet.append(["DRAG", "ratio_drag_over_brownian", "", "ratio = abs(kappa_drag_n_per_m) / abs(kappa_brownian_n_per_m)"])
            fit_sheet.append(["DRAG", "delta_n_per_m", "", "delta = kappa_drag_n_per_m - kappa_brownian_n_per_m"])
    _format_sheet(fit_sheet)
    sheet_by_key["fits"] = fit_sheet

    fit_sheet_ref = f"'{fit_sheet.title}'"
    for axis_key, sheet_key in (("x", "psd_x"), ("y", "psd_y")):
        ws = sheet_by_key.get(sheet_key)
        fit_rows = fit_param_rows.get(axis_key)
        if ws is None or fit_rows is None:
            continue
        headers = _header_map(ws)
        freq_col = headers.get("f_hz")
        data_col = headers.get("psd_um2_per_hz") or headers.get("psd_px2_per_hz")
        if freq_col is None or data_col is None:
            continue
        _add_formula_column(
            ws,
            "fit_lorentz",
            lambda row_idx, fc=freq_col, rows=fit_rows: (
                f"={fit_sheet_ref}!$C${rows['A']}/(({fit_sheet_ref}!$C${rows['fc_hz']})^2+{get_column_letter(fc)}{row_idx}^2)+{fit_sheet_ref}!$C${rows['B']}"
            ),
        )

    compare_ws = sheet_by_key.get("compare")
    if compare_ws is not None:
        headers = _header_map(compare_ws)
        brownian_col = headers.get("kappa_brownian_n_per_m")
        drag_col = headers.get("kappa_drag_n_per_m")
        if brownian_col is not None and drag_col is not None:
            _add_formula_column(
                compare_ws,
                "ratio_drag_over_brownian_excel",
                lambda row_idx, b=brownian_col, d=drag_col: (
                    f'=IF(ABS({get_column_letter(b)}{row_idx})>0,ABS({get_column_letter(d)}{row_idx})/ABS({get_column_letter(b)}{row_idx}),"")'
                ),
            )
            _add_formula_column(
                compare_ws,
                "delta_n_per_m_excel",
                lambda row_idx, b=brownian_col, d=drag_col: (
                    f"={get_column_letter(d)}{row_idx}-{get_column_letter(b)}{row_idx}"
                ),
            )

    trajectory_ws = sheet_by_key.get("trajectory")
    if trajectory_ws is not None:
        headers = _header_map(trajectory_ws)
        x_col = headers.get("x_corr_um") or headers.get("x_um") or headers.get("x_corr_px") or headers.get("x_px")
        y_col = headers.get("y_corr_um") or headers.get("y_um") or headers.get("y_corr_px") or headers.get("y_px")
        if x_col is not None and y_col is not None:
            units = "um" if "um" in str(trajectory_ws.cell(row=1, column=x_col).value).lower() else "px"
            _add_scatter_chart(
                trajectory_ws,
                x_col,
                [y_col],
                "Trajectory",
                f"x [{units}]",
                f"y [{units}]",
                "N2",
            )

    msd_ws = sheet_by_key.get("msd")
    if msd_ws is not None:
        headers = _header_map(msd_ws)
        x_col = headers.get("tau_s")
        y_cols = [headers.get(name) for name in ("msd_r_um2", "msd_r_px2", "msd_x_um2", "msd_y_um2", "msd_x_px2", "msd_y_px2")]
        y_cols = [col for col in y_cols if col is not None]
        if x_col is not None and y_cols:
            y_label = "MSD [um^2]" if headers.get("msd_r_um2") else "MSD [px^2]"
            _add_scatter_chart(msd_ws, x_col, y_cols[:3], "MSD", "tau [s]", y_label, "J2", log_x=True, log_y=True)

    for sheet_key, title in (("psd_x", "PSD X"), ("psd_y", "PSD Y")):
        ws = sheet_by_key.get(sheet_key)
        if ws is None:
            continue
        headers = _header_map(ws)
        x_col = headers.get("f_hz")
        data_cols = [headers.get("psd_um2_per_hz") or headers.get("psd_px2_per_hz")]
        fit_col = headers.get("fit_lorentz")
        if fit_col is not None:
            data_cols.append(fit_col)
        data_cols = [col for col in data_cols if col is not None]
        if x_col is not None and data_cols:
            y_label = "PSD [um^2/Hz]" if headers.get("psd_um2_per_hz") else "PSD [px^2/Hz]"
            _add_scatter_chart(ws, x_col, data_cols, title, "f [Hz]", y_label, "J2", log_x=True, log_y=True)

    for sheet_key, title in (("hist_x", "Histogram X"), ("hist_y", "Histogram Y"), ("hist_r", "Histogram R")):
        ws = sheet_by_key.get(sheet_key)
        if ws is None:
            continue
        headers = _header_map(ws)
        category_col = headers.get("bin_center_um")
        count_col = headers.get("count")
        if category_col is not None and count_col is not None:
            _add_bar_chart(ws, category_col, [count_col], title, "count", "E2")

    derived_ws = sheet_by_key.get("derived")
    if derived_ws is not None:
        headers = _header_map(derived_ws)
        category_col = headers.get("axis")
        metric_cols = [headers.get("k_pn_um"), headers.get("fc_hz")]
        metric_cols = [col for col in metric_cols if col is not None]
        if category_col is not None and metric_cols:
            _add_bar_chart(derived_ws, category_col, metric_cols, "Derived Physics Summary", "value", "N2")

    if compare_ws is not None:
        headers = _header_map(compare_ws)
        category_col = headers.get("axis") or headers.get("pair_key")
        metric_cols = [headers.get("kappa_brownian_pn_per_um"), headers.get("kappa_drag_pn_per_um")]
        metric_cols = [col for col in metric_cols if col is not None]
        if category_col is not None and metric_cols:
            _add_bar_chart(compare_ws, category_col, metric_cols, "Brownian vs Drag", "kappa [pN/um]", "N2")

    if not first_sheet_used:
        ws = wb.active
        ws.title = "README"
        ws.append(["status", "No OT CSV inputs found"])

    wb.save(out_path)
    return out_path
