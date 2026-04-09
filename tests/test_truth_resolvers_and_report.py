from __future__ import annotations

import json
import math
from pathlib import Path

import pytest
from openpyxl import load_workbook

from barakuda.core import ot_report
from barakuda.core.export_xlsx import export_ot_results_xlsx
from barakuda.core.ot_report import build_ot_item_summary, build_ot_summary_rows
from barakuda.core.truth_resolvers import (
    build_collision_safe_export_path,
    load_and_validate_timestamps_csv,
    resolve_bead_parameters_for_run,
    resolve_scale_for_run,
    resolve_timing_truth_for_run,
)


def _write(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def test_timing_truth_with_valid_timestamps(tmp_path: Path) -> None:
    video = tmp_path / "run.raw"
    video.write_bytes(b"\x00" * 16)
    ts = tmp_path / "run_timestamps.csv"
    ts.write_text(
        "frame,timestamp_s\n0,10.0\n1,10.5\n2,11.0\n",
        encoding="utf-8",
    )
    truth = resolve_timing_truth_for_run(
        video_path=video,
        frame_count=3,
        fps_hint=100.0,
        meta={"timestamps_path": str(ts)},
    )
    assert truth.timing_source == "timestamps"
    assert truth.frame_count == 3
    assert truth.elapsed_time_s == pytest.approx(1.0)
    assert truth.effective_fps == pytest.approx(2.0)


def test_timing_truth_without_timestamps_falls_back_estimated(tmp_path: Path) -> None:
    video = tmp_path / "run.raw"
    video.write_bytes(b"\x00" * 16)
    truth = resolve_timing_truth_for_run(
        video_path=video,
        frame_count=50,
        fps_hint=25.0,
        meta={},
    )
    assert truth.timing_source == "estimated"
    assert truth.elapsed_time_s == pytest.approx(2.0)
    assert truth.timestamp_validation_pass is False


def test_timing_truth_corrupted_timestamps_not_silently_accepted(tmp_path: Path) -> None:
    bad = tmp_path / "bad_timestamps.csv"
    bad.write_text(
        "frame,timestamp_s\n0,1.0\n1,0.9\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError):
        load_and_validate_timestamps_csv(bad)

    video = tmp_path / "run.raw"
    video.write_bytes(b"\x00" * 16)
    truth = resolve_timing_truth_for_run(
        video_path=video,
        frame_count=2,
        fps_hint=20.0,
        meta={"timestamps_path": str(bad)},
    )
    assert truth.timing_source == "estimated"
    assert "validation failed" in truth.timestamp_validation_message


def test_bead_provenance_explicit_values() -> None:
    b1 = resolve_bead_parameters_for_run(
        bead_diameter_um=1.0,
        bead_radius_um=None,
        bead_source="run_config",
        allow_fallback=False,
    )
    assert b1.bead_diameter_um == pytest.approx(1.0)
    assert b1.bead_radius_um == pytest.approx(0.5)
    assert b1.fallback_used is False

    b2 = resolve_bead_parameters_for_run(
        bead_diameter_um=2.0,
        bead_radius_um=None,
        bead_source="run_config",
        allow_fallback=False,
    )
    assert b2.bead_diameter_um == pytest.approx(2.0)
    assert b2.bead_radius_um == pytest.approx(1.0)


def test_bead_provenance_missing_fails_fast() -> None:
    with pytest.raises(ValueError):
        resolve_bead_parameters_for_run(
            bead_diameter_um=None,
            bead_radius_um=None,
            bead_source="missing",
            allow_fallback=False,
        )


def test_drag_summary_precedence_over_legacy(tmp_path: Path) -> None:
    run_dir = tmp_path / "analysis"
    base = "item1"
    _write(
        run_dir / "audit" / "run.json",
        {"config": {"tracking": {}, "postprocess": {}, "calibration": {}}},
    )
    _write(
        run_dir / "audit" / f"{base}_drag_summary.json",
        {"drag_force_n": 1.23, "abs_offset_um": 4.56, "kappa_pn_per_um": 7.89},
    )
    _write(
        run_dir / "audit" / f"{base}_drag.json",
        {"dragging": {"drag_force_n": 9.99}, "means_um": {"offset_um": 8.88}},
    )
    _write(
        run_dir / "audit" / f"{base}_compare.json",
        {"dragging": {"kappa_pn_per_um": 6.66}, "delta": {"ratio_drag_over_brownian": 1.1}},
    )

    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name=base,
        item_id="i1",
        source_input_path="input.raw",
        status="success",
    )
    diagnostics = summary["diagnostics"]
    metrics = summary["metrics"]
    assert diagnostics["report_source_kind"] == "drag_summary"
    assert diagnostics["drag_force_n"] == pytest.approx(1.23)
    assert diagnostics["offset_um"] == pytest.approx(4.56)
    assert metrics["kappa_drag_pn_per_um"] == pytest.approx(7.89)


def test_drag_report_reads_provenance_and_alignment_from_summary(tmp_path: Path) -> None:
    run_dir = tmp_path / "analysis"
    base = "drag_item"
    _write(
        run_dir / "audit" / "run.json",
        {"config": {"tracking": {}, "postprocess": {}, "calibration": {}}, "provenance": {}},
    )
    _write(
        run_dir / "audit" / f"{base}_drag_summary.json",
        {
            "drag_force_n": 1.23e-12,
            "abs_offset_um": 0.45,
            "kappa_pn_per_um": 6.7,
            "current_drag_input_path": "C:/drag/input.raw",
            "brownian_baseline_folder": "C:/baseline",
            "selected_calibration_path": "C:/baseline/analysis/audit/b_calibration.json",
            "selected_trajectory_path": "C:/drag/analysis/csv/drag_item_trajectory.csv",
            "selected_timestamps_path": "C:/drag/drag_item_timestamps.csv",
            "report_source_kind": "drag_summary",
            "report_source_path": "C:/drag/analysis/audit/drag_item_drag_summary.json",
            "alignment_status": "detected",
            "alignment_message": "onset detected",
            "baseline_start_s": 0.0,
            "baseline_end_s": 1.0,
            "steady_start_s": 2.0,
            "steady_end_s": 3.0,
            "actual_speed_um_s": 12.0,
            "eta_pa_s": 0.0012,
            "analysis_status": "ok",
            "physics_status": "ready",
            "warnings": ["warn1"],
        },
    )
    _write(
        run_dir / "audit" / f"{base}_alignment_diagnostics.json",
        {
            "baseline_end_s": 1.0,
            "baseline_median": 10.1,
            "baseline_mad": 0.2,
            "baseline_sigma": 0.3,
            "onset_threshold_sigma": 5.0,
            "onset_threshold_abs": 1.5,
            "onset_min_hold_s": 0.3,
            "n_baseline_samples": 100,
            "n_total_samples": 500,
            "n_frames_outside_baseline": 50,
            "candidate_onset_times_s": [2.1],
            "candidate_durations_s": [0.7],
            "failure_reason": "",
            "message": "ok",
        },
    )

    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name=base,
        item_id="i3",
        source_input_path="input.raw",
        status="success",
    )
    diagnostics = summary["diagnostics"]
    assert diagnostics["report_source_kind"] == "drag_summary"
    assert diagnostics["selected_calibration_path"] is not None
    assert diagnostics["selected_trajectory_path"] is not None
    assert diagnostics["alignment_status"] == "detected"
    assert diagnostics["alignment_baseline_sigma"] == pytest.approx(0.3)
    assert diagnostics["analysis_status"] == "ok"
    assert "warn1" in summary["warnings"]


def test_drag_summary_rows_are_drag_mode_and_include_paths(tmp_path: Path) -> None:
    run_dir = tmp_path / "analysis"
    base = "drag_item_rows"
    _write(
        run_dir / "audit" / "run.json",
        {"config": {"tracking": {}, "postprocess": {}, "calibration": {}}, "provenance": {}},
    )
    _write(
        run_dir / "audit" / "run_protocol.json",
        {
            "provenance": {
                "current_drag_input_path": "C:/drag/current_input.raw",
                "current_drag_report_path": "C:/drag/current_report.pdf",
                "brownian_baseline_folder": "C:/baseline",
                "report_source_kind": "drag_summary",
                "report_source_path": "C:/drag/analysis/audit/drag_item_rows_drag_summary.json",
            }
        },
    )
    _write(
        run_dir / "audit" / f"{base}_drag_summary.json",
        {
            "drag_force_n": 1.5e-12,
            "abs_offset_um": 0.4,
            "kappa_pn_per_um": 5.5,
            "actual_speed_um_s": 8.0,
            "eta_pa_s": 0.0011,
            "analysis_status": "ok",
            "physics_status": "ready",
            "report_source_kind": "drag_summary",
            "report_source_path": "C:/drag/analysis/audit/drag_item_rows_drag_summary.json",
        },
    )
    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name=base,
        item_id="dr1",
        source_input_path="drag.raw",
        status="success",
    )
    assert ot_report._is_drag_report(summary) is True
    rows = build_ot_summary_rows(summary)
    values = {(group, metric): value for group, metric, value, _unit, _notes in rows}
    assert values[("Drag", "Analysis mode")] == "DRAG"
    assert values[("Drag", "Drag force")] == pytest.approx(1.5e-12)
    assert values[("Drag", "Viscosity")] == pytest.approx(0.0011)
    assert values[("Drag", "Actual speed")] == pytest.approx(8.0)
    assert values[("Drag", "Absolute offset")] == pytest.approx(0.4)
    assert values[("Provenance", "Current drag input")] == "C:/drag/current_input.raw"
    assert values[("Provenance", "Current drag report path")] == "C:/drag/current_report.pdf"
    assert values[("Provenance", "Brownian baseline folder")] == "C:/baseline"
    assert values[("Provenance", "Report source kind")] == "drag_summary"


def test_drag_pdf_uses_drag_pages_not_brownian(tmp_path: Path, monkeypatch) -> None:
    run_dir = tmp_path / "analysis"
    base = "drag_pdf"
    _write(
        run_dir / "audit" / "run.json",
        {"config": {"tracking": {}, "postprocess": {}, "calibration": {}}, "provenance": {}},
    )
    _write(
        run_dir / "audit" / f"{base}_drag_summary.json",
        {
            "drag_force_n": 1.0e-12,
            "abs_offset_um": 0.2,
            "kappa_pn_per_um": 4.2,
            "actual_speed_um_s": 6.5,
            "eta_pa_s": 0.0010,
            "report_source_kind": "drag_summary",
            "report_source_path": "C:/drag/analysis/audit/drag_pdf_drag_summary.json",
        },
    )
    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name=base,
        item_id="drag_pdf_item",
        source_input_path="drag.raw",
        status="success",
    )

    class _DummyPdf:
        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc, tb):
            return False

    calls: list[str] = []

    monkeypatch.setattr("matplotlib.backends.backend_pdf.PdfPages", lambda _p: _DummyPdf())
    monkeypatch.setattr(ot_report, "_render_cover_page", lambda *args, **kwargs: calls.append("cover"))
    monkeypatch.setattr(ot_report, "_render_drag_theory_page", lambda *args, **kwargs: calls.append("drag_theory"))
    monkeypatch.setattr(ot_report, "_render_theory_page", lambda *args, **kwargs: calls.append("brownian_theory"))
    monkeypatch.setattr(ot_report, "_render_trajectory_heatmap_page", lambda *args, **kwargs: calls.append("trajectory"))
    monkeypatch.setattr(ot_report, "_render_histogram_r_and_msd_page", lambda *args, **kwargs: calls.append("hist_msd"))
    monkeypatch.setattr(ot_report, "_render_plot_pages", lambda *args, **kwargs: calls.append("plot"))
    monkeypatch.setattr(ot_report, "_render_paginated_table", lambda *args, **kwargs: calls.append("table"))
    monkeypatch.setattr(ot_report, "_render_dual_table_page", lambda *args, **kwargs: calls.append("dual"))

    ot_report.export_ot_item_pdf(tmp_path / "drag_report.pdf", summary)
    assert "drag_theory" in calls
    assert "brownian_theory" not in calls
    assert "trajectory" not in calls
    assert "hist_msd" not in calls


def test_drag_cover_status_reflects_validation_gate_fail() -> None:
    summary = {
        "status": "success",
        "item_id": "drag_item",
        "run_id": "run_1",
        "source_input_path": "C:/data/run.raw",
        "metrics": {"kappa_drag_pn_per_um": 4.0},
        "diagnostics": {
            "drag_validation_gate": "fail",
            "physics_status": "suspect_alignment_sanity",
            "physics_primary_gate": "pass",
            "detection_qc_gate": "fail",
            "final_drag_verdict": "suspect",
            "drag_force_n": 1.0e-12,
            "eta_pa_s": 0.001,
            "actual_speed_um_s": 10.0,
            "offset_um": 0.25,
            "alignment_status": "detected",
        },
    }
    left, _right = ot_report._drag_cover_rows(summary)
    status_row = next((row for row in left if row[0] == "Status"), None)
    assert status_row is not None
    assert status_row[1] == "Fail"
    combined = left + _right
    assert any(row[0] == "Primary physics gate" for row in combined)
    assert any(row[0] == "Detection QC gate" for row in combined)
    assert any(row[0] == "Final drag verdict" for row in combined)


def test_drag_report_page3_split_into_readable_blocks(tmp_path: Path, monkeypatch) -> None:
    run_dir = tmp_path / "analysis"
    base = "drag_layout"
    _write(
        run_dir / "audit" / "run.json",
        {"config": {"tracking": {}, "postprocess": {}, "calibration": {}}, "provenance": {}},
    )
    _write(
        run_dir / "audit" / f"{base}_drag_summary.json",
        {
            "drag_force_n": 1.0e-12,
            "abs_offset_um": 0.2,
            "kappa_pn_per_um": 4.2,
            "actual_speed_um_s": 6.5,
            "eta_pa_s": 0.0010,
            "report_source_kind": "drag_summary",
            "drag_validation_gate": "suspect",
            "drag_validation_reason": "alignment_sanity_fail",
            "current_drag_report_path": "C:/drag/analysis/reports/item.pdf",
        },
    )
    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name=base,
        item_id="drag_layout_item",
        source_input_path="drag.raw",
        status="success",
    )
    class _DummyPdf:
        def __enter__(self):
            return self
        def __exit__(self, exc_type, exc, tb):
            return False

    page_titles: list[str] = []
    monkeypatch.setattr("matplotlib.backends.backend_pdf.PdfPages", lambda _p: _DummyPdf())
    monkeypatch.setattr(ot_report, "_render_cover_page", lambda *args, **kwargs: None)
    monkeypatch.setattr(ot_report, "_render_drag_theory_page", lambda *args, **kwargs: None)
    monkeypatch.setattr(
        ot_report,
        "_render_paginated_table",
        lambda _pdf, title, *_args, **_kwargs: page_titles.append(title),
    )
    monkeypatch.setattr(ot_report, "_render_plot_pages", lambda *args, **kwargs: None)

    ot_report.export_ot_item_pdf(tmp_path / "drag_layout.pdf", summary)
    assert "Drag Provenance" in page_titles
    assert "Drag Timing and Alignment" in page_titles
    assert "Drag QC and Confidence" in page_titles
    assert "Drag Warnings" in page_titles


def test_drag_report_trace_page_prefers_relative_time_axis(tmp_path: Path) -> None:
    trace = tmp_path / "trace.csv"
    trace.write_text("video_time_rel_s,video_time_s,axis_px\n0.0,10.0,1.0\n", encoding="utf-8")
    entry = ot_report._build_plot_entry(
        "Annotated drag trace",
        str(trace),
        ("video_time_rel_s", "video_time_s", "stage_time_aligned_s"),
        ("axis_px",),
        "Time from video start [s]",
        "Axis position [px]",
        False,
        False,
    )
    assert entry is not None
    assert entry[2][0] == "video_time_rel_s"
    assert entry[4] == "Time from video start [s]"


def test_current_drag_report_path_prefers_pdf_over_summary_json(tmp_path: Path) -> None:
    run_dir = tmp_path / "analysis"
    base = "drag_pdf_path"
    _write(
        run_dir / "audit" / "run_protocol.json",
        {
            "provenance": {
                "current_drag_report_path": "C:/drag/reports/final_report.pdf",
                "current_drag_summary_json_path": "C:/drag/audit/drag_pdf_path_drag_summary.json",
            }
        },
    )
    _write(
        run_dir / "audit" / "run.json",
        {"config": {"tracking": {}, "postprocess": {}, "calibration": {}}, "provenance": {}},
    )
    _write(
        run_dir / "audit" / f"{base}_drag_summary.json",
        {
            "drag_force_n": 1.0e-12,
            "abs_offset_um": 0.2,
            "kappa_pn_per_um": 4.2,
            "actual_speed_um_s": 6.5,
            "eta_pa_s": 0.0010,
            "report_source_kind": "drag_summary",
        },
    )
    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name=base,
        item_id="drag_pdf_item",
        source_input_path="drag.raw",
        status="success",
    )
    assert summary["diagnostics"].get("current_drag_report_path") == "C:/drag/reports/final_report.pdf"


def test_trace_page_renders_primary_and_qc_markers(tmp_path: Path, monkeypatch) -> None:
    trace = tmp_path / "trace.csv"
    trace.write_text(
        "video_time_rel_s,video_time_s,axis_px\n0.0,10.0,1.0\n1.0,11.0,1.1\n2.0,12.0,1.2\n",
        encoding="utf-8",
    )
    summary = {
        "diagnostics": {
            "baseline_start_s": 0.0,
            "baseline_end_s": 0.5,
            "steady_start_s": 1.2,
            "steady_end_s": 1.8,
            "expected_stage_start_video_s": 11.0,
            "expected_stage_stop_video_s": 12.0,
            "detected_stage_start_video_s": 11.4,
            "detected_stage_stop_video_s": 12.4,
            "t_first_s": 10.0,
        }
    }
    labels: list[str] = []
    import matplotlib.axes

    original_axvline = matplotlib.axes.Axes.axvline

    def _capture_axvline(self, x=0, *args, **kwargs):
        label = kwargs.get("label")
        if isinstance(label, str):
            labels.append(label)
        return original_axvline(self, x, *args, **kwargs)

    monkeypatch.setattr(matplotlib.axes.Axes, "axvline", _capture_axvline)

    class _DummyPdf:
        def savefig(self, _fig):
            return None

    ot_report._render_drag_trace_page(_DummyPdf(), trace, summary)
    assert "expected stage start" in labels
    assert "expected stage stop" in labels
    assert "detected onset (QC)" in labels
    assert "detected stop (QC)" in labels


def test_xlsx_drag_sheet_exposes_gate_and_verdict_fields(tmp_path: Path) -> None:
    base = "drag_xlsx"
    run_dir = tmp_path / "analysis"
    output_dir = run_dir / "results"
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_dir = run_dir / "csv"
    audit_dir = run_dir / "audit"
    csv_dir.mkdir(parents=True, exist_ok=True)
    audit_dir.mkdir(parents=True, exist_ok=True)

    traj = csv_dir / f"{base}_trajectory.csv"
    traj.write_text("frame,t_s,x_px,y_px\n0,0.0,1.0,1.0\n1,0.1,1.1,1.1\n", encoding="utf-8")
    _write(
        audit_dir / f"{base}_drag_summary.json",
        {
            "drag_force_n": 1.0e-12,
            "eta_pa_s": 0.0012,
            "physics_primary_gate": "suspect",
            "detection_qc_gate": "fail",
            "final_drag_verdict": "suspect",
            "stage_anchor_confidence": "high",
            "detected_onset_consistency_flag": False,
            "baseline_strategy_difference_ratio": 1.01,
            "speed_consistency_error_pct": 0.74,
            "measured_offset_um": 0.028,
            "expected_offset_if_eta_1mPas_um": 0.023,
        },
    )
    _write(
        audit_dir / "run_protocol.json",
        {"provenance": {"current_drag_report_path": "C:/drag/reports/drag_xlsx.pdf"}},
    )

    xlsx_path = export_ot_results_xlsx(
        output_dir=output_dir,
        base_name=base,
        trajectory_csv_path=traj,
    )
    wb = load_workbook(xlsx_path)
    ws = wb["DRAG"]
    pairs = {}
    for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        key, value = row
        if key is None:
            continue
        pairs[str(key)] = value
    assert pairs.get("drag.physics_primary_gate") == "suspect"
    assert pairs.get("drag.detection_qc_gate") == "fail"
    assert pairs.get("drag.final_drag_verdict") == "suspect"
    assert pairs.get("drag.stage_anchor_confidence") == "high"
    assert pairs.get("drag.detected_onset_consistency_flag") in {"False", False}
    assert pairs.get("drag.baseline_strategy_difference_ratio") is not None
    assert pairs.get("drag.speed_consistency_error_pct") is not None
    assert pairs.get("drag.measured_offset_um") is not None
    assert pairs.get("drag.expected_offset_if_eta_1mPas_um") is not None
    assert pairs.get("provenance.current_drag_report_path") == "C:/drag/reports/drag_xlsx.pdf"


def test_drag_summary_parses_physics_confidence_fields(tmp_path: Path) -> None:
    run_dir = tmp_path / "analysis"
    base = "drag_conf"
    _write(
        run_dir / "audit" / "run.json",
        {"config": {"tracking": {}, "postprocess": {}, "calibration": {}}, "provenance": {}},
    )
    _write(
        run_dir / "audit" / f"{base}_drag_summary.json",
        {
            "drag_force_n": 1.0e-12,
            "abs_offset_um": 0.2,
            "kappa_pn_per_um": 4.2,
            "actual_speed_um_s": 5.5,
            "eta_pa_s": 0.001,
            "report_source_kind": "drag_summary",
            "drag_physics_confidence": "medium",
            "drag_physics_warning": "baseline sensitivity",
            "baseline_robustness_flag": "weak",
            "onset_robustness_flag": "weak",
            "kinematics_robustness_flag": "legacy_units_risky",
            "eta_current_windows": 0.001,
            "eta_alt_baseline": 0.0016,
            "offset_current_windows_um": 0.2,
            "offset_alt_baseline_um": 0.31,
            "stage_speed_from_trace_um_s": 80.1,
            "stage_speed_relative_diff": 0.02,
            "stage_speed_consistent": True,
            "baseline_strategy_primary": "near_onset_baseline",
            "baseline_strategy_alt": "long_premotion_baseline_reference",
            "offset_primary_um": 0.2,
            "offset_alt_um": 0.31,
            "eta_primary_pa_s": 0.001,
            "eta_alt_pa_s": 0.0016,
            "baseline_strategy_difference_ratio": 1.6,
            "onset_confidence_class": "medium",
            "speed_stage_json": 82.0,
            "speed_trace_derived": 80.1,
            "speed_used_for_physics": 82.0,
            "speed_consistency_error_pct": 2.3,
            "drag_validation_gate": "suspect",
            "drag_validation_reason": "baseline_suspect,kinematics_suspect",
            "alignment_sanity_flag": False,
            "alignment_sanity_message": "detected_stop_after_video_end",
            "expected_stage_start_video_s": 1.2,
            "expected_stage_stop_video_s": 9.8,
            "detected_stage_start_video_s": 1.25,
            "detected_stage_stop_video_s": 10.3,
        },
    )
    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name=base,
        item_id="c1",
        source_input_path="drag.raw",
        status="success",
    )
    diagnostics = summary["diagnostics"]
    assert diagnostics["drag_physics_confidence"] == "medium"
    assert diagnostics["baseline_robustness_flag"] == "weak"
    assert diagnostics["eta_alt_baseline"] == pytest.approx(0.0016)
    assert diagnostics["stage_speed_consistent"] is True
    rows = build_ot_summary_rows(summary)
    row_map = {(g, m): v for g, m, v, _u, _n in rows}
    assert row_map[("Drag", "Drag validation gate")] == "suspect"
    assert row_map[("Drag", "Baseline strategy difference ratio")] == pytest.approx(1.6)
    assert row_map[("Drag", "Alignment sanity flag")] is False


def test_dropped_frames_semantics_kept_separate(tmp_path: Path) -> None:
    run_dir = tmp_path / "analysis"
    base = "item2"
    _write(
        run_dir / "audit" / "run.json",
        {
            "config": {
                "tracking": {},
                "postprocess": {},
                "calibration": {},
                "acquisition": {"camera_dropped_frames": 3},
            }
        },
    )
    _write(
        run_dir / "audit" / f"{base}_postprocess.json",
        {"summary": {"qc": {"lost_frames": 11, "lost_fraction": 0.11}}},
    )
    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name=base,
        item_id="i2",
        source_input_path="input.raw",
        status="success",
    )
    diagnostics = summary["diagnostics"]
    assert diagnostics["camera_dropped_frames"] == pytest.approx(3.0)
    assert diagnostics["tracking_lost_frames"] == pytest.approx(11.0)
    assert diagnostics["lost_fraction"] == pytest.approx(0.11)


def test_brownian_conditions_qc_do_not_include_drag_na_ballast() -> None:
    summary = {
        "metrics": {
            "mode": "BROWNIAN",
            "fps": 1000.0,
            "um_per_px": 0.06042,
            "scale_source": "dataset",
            "temperature_c": 25.0,
            "bead_diameter_um": 2.0,
            "bead_radius_um": 1.0,
            "bead_source": "run_config",
        },
        "diagnostics": {
            "effective_fps": 999.0,
            "elapsed_time_s": 10.0,
            "timing_source": "timestamps",
            "selected_calibration_path": "C:/run/audit/c.json",
            "selected_trajectory_path": "C:/run/csv/t.csv",
            "selected_timestamps_path": "C:/run/raw/ts.csv",
            "report_source_kind": "brownian",
            "report_source_path": "C:/run/audit/run.json",
            "lost_fraction": 0.01,
            "camera_dropped_frames": 0,
            "tracking_lost_frames": 2,
            "timestamp_validation_pass": True,
            "timestamp_validation_message": "ok",
            "warning_count": 0,
        },
        "warnings": [],
    }
    cond_rows = ot_report._conditions_rows(summary)
    cond_labels = [r[0] for r in cond_rows]
    assert "Current drag input" not in cond_labels
    assert "Brownian baseline folder" not in cond_labels

    qc_rows = ot_report._qc_rows(summary)
    qc_labels = [r[0] for r in qc_rows]
    assert "Drag force [N]" not in qc_labels
    assert "Alignment offset [s]" not in qc_labels


def test_brownian_preview_selection_is_representative_and_limited() -> None:
    previews = [Path(f"C:/run/raw/video_preview_{i:02d}.png") for i in range(1, 11)]
    selected = ot_report._select_representative_preview_paths(previews, max_images=6)
    assert len(selected) == 6
    assert selected[0].name == "video_preview_01.png"
    assert selected[-1].name == "video_preview_10.png"


def test_brownian_identity_rows_use_acquisition_and_processed_times(tmp_path: Path) -> None:
    item_root = tmp_path / "items" / "Gly20_brown_rep01"
    analysis_dir = item_root / "analysis"
    raw_dir = item_root / "raw"
    analysis_dir.mkdir(parents=True, exist_ok=True)
    raw_dir.mkdir(parents=True, exist_ok=True)
    (raw_dir / "video_meta.json").write_text(
        json.dumps({"timestamp": "2026-04-08T11:23:27"}, indent=2),
        encoding="utf-8",
    )
    summary = {
        "item_id": "Gly20_brown_rep01",
        "status": "success",
        "analysis_dir": str(analysis_dir),
        "created_at": "2026-04-08T16:50:15",
    }
    rows = ot_report._identity_rows(summary)
    labels = {row[0]: row[1] for row in rows}
    assert labels["Acquisition date/time"] == "2026-04-08T11:23:27"
    assert labels["Processed / report generated"] == "2026-04-08T16:50:15"
    assert "Acquisition source" not in labels


def test_scientific_notation_uses_multiplication_sign() -> None:
    s = ot_report._fmt_scientific_text(2.46e-3, sig_figs=3)
    assert "× 10" in s


def test_relaxation_time_from_fc_hz() -> None:
    fc = 57.77372569470335
    tau, se = ot_report._relaxation_time_from_fc_hz(fc, 2.0095666970207935e-05)
    assert tau == pytest.approx(1.0 / (2.0 * math.pi * fc))
    assert se is not None
    assert se == pytest.approx(2.0095666970207935e-05 / (2.0 * math.pi * fc * fc))


def test_brownian_summary_includes_uncertainties_tau_anisotropy_qc(tmp_path: Path) -> None:
    run_dir = tmp_path / "analysis"
    base = "brown_run"
    audit = run_dir / "audit"
    audit.mkdir(parents=True)
    _write(
        audit / "run.json",
        {
            "config": {
                "tracking": {
                    "timing": {
                        "timestamp_validation_pass": True,
                        "timestamp_validation_message": "ok",
                    }
                },
                "postprocess": {"physics_mode": "BROWNIAN"},
                "calibration": {},
            }
        },
    )
    _write(
        audit / f"{base}_postprocess.json",
        {
            "summary": {
                "qc": {"lost_frames": 0, "lost_fraction": 0.0},
                "physics": {
                    "lorentz_fit": {
                        "x": {"rmse": 8.76e-6, "A": 1.112},
                        "y": {"rmse": 7.05e-6, "A": 1.375},
                    }
                },
            }
        },
    )
    _write(
        audit / f"{base}_calibration.json",
        {
            "kappa": {
                "kappa_x_pn_per_um": 18.0,
                "kappa_y_pn_per_um": 17.0,
                "kappa_x_pn_per_um_se": 0.07,
                "kappa_y_pn_per_um_se": 0.07,
            },
            "viscosity": {"eta_mean_pa_s": 0.002, "eta_mean_pa_s_se": 1e-5},
            "diffusion": {"D_m2_s": 1e-13, "D_m2_s_se": 1e-15},
            "diagnostics": {
                "fc_x_hz": 57.0,
                "fc_y_hz": 64.0,
                "fc_x_hz_se": 0.02,
                "fc_y_hz_se": 0.02,
            },
            "kappa_unit_check": {"pass": True},
            "anisotropy": {"pass": True},
        },
    )
    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name=base,
        item_id=base,
        source_input_path="x.raw",
        status="success",
    )
    d = summary["diagnostics"]
    assert d["tau_x_s"] * 2.0 * math.pi * d["fc_x_hz"] == pytest.approx(1.0)
    assert d["tau_y_s"] * 2.0 * math.pi * d["fc_y_hz"] == pytest.approx(1.0)
    assert d["trap_kappa_ratio_xy"] == pytest.approx(18.0 / 17.0)
    assert d["trap_fc_ratio_xy"] == pytest.approx(57.0 / 64.0)
    assert d["brownian_qc"]["timing"] == "pass"
    assert d["brownian_qc"]["overall"] == "usable"
    key_rows = ot_report._key_result_rows(summary)
    text_blob = " ".join(f"{a} {b}" for a, b in key_rows)
    assert "±" in text_blob or "×" in text_blob
    low = text_blob.lower()
    assert "rms" not in low
    assert "variance" not in low
    labels = [a for a, _ in key_rows]
    assert any("Trap anisotropy" in lab for lab in labels)
    assert any("Corner-frequency ratio" in lab for lab in labels)
    assert any(lab.startswith("QC:") for lab in labels)
    assert "Diffusion coefficient" in labels
    diff_row = next(r for r in key_rows if r[0] == "Diffusion coefficient")
    assert ("×" in diff_row[1]) or ("e-" in diff_row[1].lower()) or ("µm²/s" in diff_row[1])
    qc = ot_report._qc_rows(summary)
    qc_labels = [r[0] for r in qc]
    assert "QC rules (summary)" in qc_labels
    assert "QC: timing" not in qc_labels
    assert "QC: tracking" not in qc_labels
    assert "QC: PSD fit" not in qc_labels
    assert "Overall Brownian QC" not in qc_labels


def test_brownian_cover_key_results_use_readable_display_units() -> None:
    summary = {
        "metrics": {
            "mode": "BROWNIAN",
            "eta_mean_pa_s": 0.00123,
            "eta_mean_pa_s_se": 1.0e-5,
            "D_m2_s": 2.5e-13,
            "D_m2_s_se": 1.0e-14,
            "kappa_x_pn_per_um": 18.0,
            "kappa_x_pn_per_um_se": 0.07,
            "kappa_y_pn_per_um": 17.0,
            "kappa_y_pn_per_um_se": 0.07,
        },
        "diagnostics": {
            "fc_x_hz": 57.0,
            "fc_x_hz_se": 0.02,
            "fc_y_hz": 64.0,
            "fc_y_hz_se": 0.02,
            "tau_x_s": 1.0 / (2.0 * math.pi * 57.0),
            "tau_x_s_se": 1.0e-5,
            "tau_y_s": 1.0 / (2.0 * math.pi * 64.0),
            "tau_y_s_se": 1.0e-5,
            "trap_kappa_ratio_xy": 18.0 / 17.0,
            "trap_fc_ratio_xy": 57.0 / 64.0,
            "brownian_qc": {"timing": "pass", "tracking": "pass", "psd_fit": "pass", "overall": "usable"},
        },
    }
    rows = ot_report._key_result_rows(summary)
    value_by_label = {label: value for label, value in rows}
    assert "mPa" in value_by_label["Mean viscosity"]
    assert "µm²/s" in value_by_label["Diffusion coefficient"]
    assert "ms" in value_by_label["Relaxation time X"]
    assert "ms" in value_by_label["Relaxation time Y"]
    assert "Trap anisotropy (κx/κy)" in value_by_label
    assert "QC: timing" in value_by_label
    assert "Overall Brownian QC" in value_by_label


def test_brownian_qc_verdict_stays_on_cover_not_page3() -> None:
    summary = {
        "metrics": {"mode": "BROWNIAN"},
        "diagnostics": {
            "brownian_qc": {
                "timing": "pass",
                "tracking": "pass",
                "psd_fit": "warning",
                "overall": "caution",
                "notes": "summary notes",
            }
        },
    }
    cover_rows = ot_report._key_result_rows(summary)
    cover_labels = [a for a, _ in cover_rows]
    assert "QC: timing" in cover_labels
    assert "QC: tracking" in cover_labels
    assert "QC: PSD fit" in cover_labels
    assert "Overall Brownian QC" in cover_labels

    page3_rows = ot_report._qc_rows(summary)
    page3_labels = [a for a, _ in page3_rows]
    assert "QC: timing" not in page3_labels
    assert "QC: tracking" not in page3_labels
    assert "QC: PSD fit" not in page3_labels
    assert "Overall Brownian QC" not in page3_labels
    assert "QC rules (summary)" in page3_labels


def test_relaxation_time_cover_format_is_shorter_in_ms() -> None:
    summary = {
        "metrics": {"mode": "BROWNIAN"},
        "diagnostics": {
            "tau_x_s": 0.0027561234,
            "tau_x_s_se": 0.000000321,
            "tau_y_s": 0.0031049876,
            "tau_y_s_se": 0.000000456,
        },
    }
    rows = ot_report._key_result_rows(summary)
    values = {label: value for label, value in rows}
    tx = values["Relaxation time X"]
    ty = values["Relaxation time Y"]
    assert "ms" in tx and "ms" in ty
    assert tx.count(".") <= 2
    assert ty.count(".") <= 2


def test_cover_single_table_rows_drop_item_and_source_when_title_has_item() -> None:
    left_rows = [
        ["Item", "Gly20_brown_rep01"],
        ["Status", "Success"],
        ["Acquisition date/time", "2026-04-08T11:23:27"],
        ["Acquisition source", "video_meta.timestamp"],
        ["Processed / report generated", "2026-04-08T16:50:15"],
    ]
    right_rows = [
        ["Mean viscosity", "1.230 ± 0.010 mPa·s"],
        ["QC: timing", "pass"],
        ["Overall Brownian QC", "usable"],
    ]
    rows = ot_report._build_single_cover_table_rows(
        "Item Report: Gly20_brown_rep01",
        left_rows,
        right_rows,
    )
    labels = [r[0] for r in rows]
    assert "Item" not in labels
    assert "Acquisition source" not in labels
    assert "Identity / timing" in labels
    assert "Main results" in labels
    assert "Brownian QC" in labels


def test_render_cover_page_uses_single_table_not_multi_cards(monkeypatch) -> None:
    class _DummyPdf:
        def savefig(self, _fig):
            return None

    called = {"cards": 0}

    def _forbid_card(*_args, **_kwargs):
        called["cards"] += 1
        raise AssertionError("multi-card cover layout should not be used")

    monkeypatch.setattr(ot_report, "_render_cover_card", _forbid_card)
    ot_report._render_cover_page(
        _DummyPdf(),
        title="Item Report: Gly20_brown_rep01",
        subtitle="Item-level summary",
        left_rows=[
            ["Item", "Gly20_brown_rep01"],
            ["Status", "Success"],
            ["Acquisition date/time", "2026-04-08T11:23:27"],
            ["Processed / report generated", "2026-04-08T16:50:15"],
        ],
        right_rows=[
            ["Mean viscosity", "1.230 ± 0.010 mPa·s"],
            ["QC: timing", "pass"],
            ["Overall Brownian QC", "usable"],
        ],
    )
    assert called["cards"] == 0


def test_real_item_preview_v6_generation_smoke() -> None:
    item_root = Path(r"C:/Work/BARAKUDA_FULL/runs/ot/2026-04-08_164655/items/Gly20_brown_rep01")
    run_dir = item_root / "analysis"
    out_pdf = run_dir / "results" / "2026-04-09-OT-Gly20_brown_rep01-report-polished-preview-v6.pdf"
    if not run_dir.is_dir():
        pytest.skip("Real item folder is not available in this environment.")
    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name="Gly20_brown_rep01",
        item_id="Gly20_brown_rep01",
        source_input_path=str(item_root / "raw" / "Gly20_brown_rep01.raw"),
        status="success",
    )
    summary["analysis_dir"] = str(run_dir)
    ot_report.export_ot_item_pdf(out_pdf, summary)
    assert out_pdf.is_file()
    assert out_pdf.stat().st_size > 1000


def test_real_item_preview_v7_generation_smoke() -> None:
    item_root = Path(r"C:/Work/BARAKUDA_FULL/runs/ot/2026-04-08_164655/items/Gly20_brown_rep01")
    run_dir = item_root / "analysis"
    out_pdf = run_dir / "results" / "2026-04-09-OT-Gly20_brown_rep01-report-polished-preview-v7.pdf"
    if not run_dir.is_dir():
        pytest.skip("Real item folder is not available in this environment.")
    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name="Gly20_brown_rep01",
        item_id="Gly20_brown_rep01",
        source_input_path=str(item_root / "raw" / "Gly20_brown_rep01.raw"),
        status="success",
    )
    summary["analysis_dir"] = str(run_dir)
    ot_report.export_ot_item_pdf(out_pdf, summary)
    assert out_pdf.is_file()
    assert out_pdf.stat().st_size > 1000


def test_histogram_and_msd_page_emphasizes_msd_log_curve(tmp_path: Path, monkeypatch) -> None:
    hist = tmp_path / "hist_r.csv"
    hist.write_text("bin_center_um,count\n0.1,10\n0.2,15\n0.3,8\n", encoding="utf-8")
    msd = tmp_path / "msd.csv"
    msd.write_text(
        "tau_s,msd_r_um2\n0.001,0.002\n0.002,0.004\n0.004,0.009\n0.008,0.018\n",
        encoding="utf-8",
    )

    class _DummyPdf:
        def savefig(self, _fig):
            return None

    import matplotlib.axes

    called = {"loglog": 0}
    original_loglog = matplotlib.axes.Axes.loglog

    def _capture_loglog(self, *args, **kwargs):
        called["loglog"] += 1
        assert kwargs.get("linewidth", 0) >= 1.9
        return original_loglog(self, *args, **kwargs)

    monkeypatch.setattr(matplotlib.axes.Axes, "loglog", _capture_loglog)
    ot_report._render_histogram_r_and_msd_page(_DummyPdf(), hist, msd)
    assert called["loglog"] >= 1


def test_graph_pages_emit_short_captions(tmp_path: Path, monkeypatch) -> None:
    caps: list[str] = []
    monkeypatch.setattr(ot_report, "_add_figure_caption", lambda _fig, text, **_kwargs: caps.append(str(text)))

    class _DummyPdf:
        def savefig(self, _fig):
            return None

    # Preview frames page
    import numpy as np
    import matplotlib.pyplot as plt

    img = tmp_path / "p1.png"
    plt.imsave(img, np.zeros((8, 8, 3), dtype=np.uint8))
    ot_report._render_preview_grid_page(_DummyPdf(), [img])

    # Trajectory page
    traj = tmp_path / "traj.csv"
    traj.write_text("x_corr_um,y_corr_um\n0.0,0.0\n0.1,0.05\n0.2,-0.03\n", encoding="utf-8")
    ot_report._render_trajectory_heatmap_page(_DummyPdf(), traj)

    # Histogram+MSD page
    hist = tmp_path / "hist_r.csv"
    hist.write_text("bin_center_um,count\n0.1,10\n0.2,12\n", encoding="utf-8")
    msd = tmp_path / "msd.csv"
    msd.write_text("tau_s,msd_r_um2\n0.001,0.002\n0.002,0.004\n", encoding="utf-8")
    ot_report._render_histogram_r_and_msd_page(_DummyPdf(), hist, msd)

    assert any("Representative frames" in c for c in caps)
    assert any("Heat map and marginal histograms" in c for c in caps)
    assert any("Histogram R summarizes" in c for c in caps)


def test_psd_plot_page_draws_measured_and_fit_overlay(tmp_path: Path, monkeypatch) -> None:
    csv_dir = tmp_path / "analysis" / "csv"
    audit_dir = tmp_path / "analysis" / "audit"
    csv_dir.mkdir(parents=True)
    audit_dir.mkdir(parents=True)

    A, B, fc = 1.5, 0.01, 3.0
    psd_x = csv_dir / "demo_psd_x.csv"
    freqs = [float(i) for i in range(1, 65)]
    vals = [(A / (fc * fc + f * f)) + B for f in freqs]
    lines = ["f_hz,psd_um2_per_hz"] + [f"{f},{v}" for f, v in zip(freqs, vals)]
    psd_x.write_text("\n".join(lines) + "\n", encoding="utf-8")
    (audit_dir / "demo_psd_fit.json").write_text(
        json.dumps({"fit_x": {"A": A, "B": B, "fc_hz": fc, "fmin_hz": 1.0, "fmax_hz": 64.0}}, indent=2),
        encoding="utf-8",
    )

    calls: list[str] = []
    fc_lines = {"count": 0}
    fc_labels: list[str] = []
    import matplotlib.axes

    original_plot = matplotlib.axes.Axes.plot
    original_axvline = matplotlib.axes.Axes.axvline
    original_text = matplotlib.axes.Axes.text

    def _capture_plot(self, *args, **kwargs):
        label = kwargs.get("label")
        if isinstance(label, str):
            calls.append(label)
        return original_plot(self, *args, **kwargs)

    def _capture_axvline(self, x=0, *args, **kwargs):
        fc_lines["count"] += 1
        return original_axvline(self, x, *args, **kwargs)

    def _capture_text(self, x, y, s, *args, **kwargs):
        if isinstance(s, str) and "fc =" in s:
            fc_labels.append(s)
        return original_text(self, x, y, s, *args, **kwargs)

    monkeypatch.setattr(matplotlib.axes.Axes, "plot", _capture_plot)
    monkeypatch.setattr(matplotlib.axes.Axes, "axvline", _capture_axvline)
    monkeypatch.setattr(matplotlib.axes.Axes, "text", _capture_text)

    class _DummyPdf:
        def savefig(self, _fig):
            return None

    entry = ot_report._build_plot_entry(
        "PSD X",
        str(psd_x),
        ("f_hz",),
        ("psd_um2_per_hz",),
        "Frequency [Hz]",
        "PSD",
        True,
        True,
    )
    assert entry is not None
    ot_report._render_plot_pages(_DummyPdf(), "Power Spectral Density", [entry], layout="vertical", items_per_page=2)
    assert "Measured PSD" in calls
    assert "Fitted PSD" in calls
    assert fc_lines["count"] >= 1
    assert any("Hz" in t for t in fc_labels)


def test_psd_overlay_fallback_when_inconsistent_fit(tmp_path: Path, monkeypatch) -> None:
    csv_dir = tmp_path / "analysis" / "csv"
    audit_dir = tmp_path / "analysis" / "audit"
    csv_dir.mkdir(parents=True)
    audit_dir.mkdir(parents=True)

    psd_x = csv_dir / "bad_psd_x.csv"
    psd_x.write_text(
        "f_hz,psd_um2_per_hz\n1.0,1.2\n2.0,0.9\n4.0,0.6\n8.0,0.3\n16.0,0.2\n32.0,0.1\n",
        encoding="utf-8",
    )
    (audit_dir / "bad_psd_fit.json").write_text(
        json.dumps({"fit_x": {"A": 1e-10, "B": 1e-12, "fc_hz": 50.0, "fmin_hz": 1.0, "fmax_hz": 32.0}}, indent=2),
        encoding="utf-8",
    )

    calls: list[str] = []
    import matplotlib.axes

    original_plot = matplotlib.axes.Axes.plot

    def _capture_plot(self, *args, **kwargs):
        label = kwargs.get("label")
        if isinstance(label, str):
            calls.append(label)
        return original_plot(self, *args, **kwargs)

    monkeypatch.setattr(matplotlib.axes.Axes, "plot", _capture_plot)

    class _DummyPdf:
        def savefig(self, _fig):
            return None

    entry = ot_report._build_plot_entry(
        "PSD X",
        str(psd_x),
        ("f_hz",),
        ("psd_um2_per_hz",),
        "Frequency [Hz]",
        "PSD",
        True,
        True,
    )
    assert entry is not None
    ot_report._render_plot_pages(_DummyPdf(), "Power Spectral Density", [entry], layout="vertical", items_per_page=2)
    assert "Measured PSD" in calls
    assert "Fitted PSD" not in calls


def test_qc_rules_summary_is_compact_for_table_layout() -> None:
    summary = {
        "metrics": {"mode": "BROWNIAN"},
        "diagnostics": {
            "brownian_qc": {
                "notes": "very long original note should be compacted in table",
            }
        },
    }
    rows = ot_report._qc_rows(summary)
    qc_row = next((r for r in rows if r[0] == "QC rules (summary)"), None)
    assert qc_row is not None
    text = qc_row[1]
    assert "Timing:" in text and "Tracking:" in text and "PSD:" in text
    assert len(str(text)) < 220


def test_style_table_reduces_qc_rules_value_cell_padding() -> None:
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots()
    ax.axis("off")
    table = ax.table(
        cellText=[
            ["QC rules (summary)", "Timing: ... Tracking: ... PSD: ..."],
            ["Warnings", "0"],
        ],
        colLabels=["Field", "Value"],
        cellLoc="left",
        colLoc="left",
        bbox=[0.0, 0.0, 1.0, 1.0],
    )
    ot_report._style_table(table)
    # header is row 0; first body row is row 1.
    assert table[(1, 1)].PAD <= 0.08
    plt.close(fig)


def test_safe_psd_overlay_applies_unit_scale_when_needed(tmp_path: Path) -> None:
    analysis = tmp_path / "analysis"
    csv_dir = analysis / "csv"
    audit_dir = analysis / "audit"
    csv_dir.mkdir(parents=True)
    audit_dir.mkdir(parents=True)
    um_per_px = 0.1
    (audit_dir / "run.json").write_text(
        json.dumps({"config": {"calibration": {"um_per_px": um_per_px}}}, indent=2),
        encoding="utf-8",
    )
    A, B, fc = 1.0, 0.02, 5.0
    (audit_dir / "u_psd_fit.json").write_text(
        json.dumps({"fit_x": {"A": A, "B": B, "fc_hz": fc, "fmin_hz": 1.0, "fmax_hz": 64.0}}, indent=2),
        encoding="utf-8",
    )
    psd_path = csv_dir / "u_psd_x.csv"
    xs = [float(i) for i in range(1, 65)]
    # Measured in um^2/Hz, while fit parameters are effectively px^2/Hz.
    ys = [((A / (fc * fc + x * x)) + B) * (um_per_px**2) for x in xs]
    lines = ["f_hz,psd_um2_per_hz"] + [f"{x},{y}" for x, y in zip(xs, ys)]
    psd_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    overlay = ot_report._safe_psd_fit_overlay(
        label="PSD X",
        csv_path=psd_path,
        x_values=xs,
        y_values=ys,
        y_column_name="psd_um2_per_hz",
    )
    assert overlay is not None
    ox, oy, fc_fit = overlay
    assert len(ox) >= 20
    assert len(ox) == len(oy)
    assert fc_fit > 0


def test_msd_page_does_not_draw_fake_fit_overlay(tmp_path: Path, monkeypatch) -> None:
    hist = tmp_path / "hist_r.csv"
    hist.write_text("bin_center_um,count\n0.1,10\n0.2,12\n", encoding="utf-8")
    msd = tmp_path / "msd.csv"
    msd.write_text("tau_s,msd_r_um2\n0.001,0.002\n0.002,0.004\n0.004,0.009\n", encoding="utf-8")

    labels: list[str] = []
    import matplotlib.axes

    original_plot = matplotlib.axes.Axes.plot

    def _capture_plot(self, *args, **kwargs):
        label = kwargs.get("label")
        if isinstance(label, str):
            labels.append(label)
        return original_plot(self, *args, **kwargs)

    monkeypatch.setattr(matplotlib.axes.Axes, "plot", _capture_plot)

    class _DummyPdf:
        def savefig(self, _fig):
            return None

    ot_report._render_histogram_r_and_msd_page(_DummyPdf(), hist, msd)
    assert all("fit" not in str(l).lower() for l in labels)


def test_real_item_preview_v8_generation_smoke() -> None:
    item_root = Path(r"C:/Work/BARAKUDA_FULL/runs/ot/2026-04-08_164655/items/Gly20_brown_rep01")
    run_dir = item_root / "analysis"
    out_pdf = run_dir / "results" / "2026-04-09-OT-Gly20_brown_rep01-report-polished-preview-v8.pdf"
    if not run_dir.is_dir():
        pytest.skip("Real item folder is not available in this environment.")
    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name="Gly20_brown_rep01",
        item_id="Gly20_brown_rep01",
        source_input_path=str(item_root / "raw" / "Gly20_brown_rep01.raw"),
        status="success",
    )
    summary["analysis_dir"] = str(run_dir)
    ot_report.export_ot_item_pdf(out_pdf, summary)
    assert out_pdf.is_file()
    assert out_pdf.stat().st_size > 1000


def test_real_item_preview_v9_generation_smoke() -> None:
    item_root = Path(r"C:/Work/BARAKUDA_FULL/runs/ot/2026-04-08_164655/items/Gly20_brown_rep01")
    run_dir = item_root / "analysis"
    out_pdf = run_dir / "results" / "2026-04-09-OT-Gly20_brown_rep01-report-polished-preview-v9.pdf"
    if not run_dir.is_dir():
        pytest.skip("Real item folder is not available in this environment.")
    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name="Gly20_brown_rep01",
        item_id="Gly20_brown_rep01",
        source_input_path=str(item_root / "raw" / "Gly20_brown_rep01.raw"),
        status="success",
    )
    summary["analysis_dir"] = str(run_dir)
    ot_report.export_ot_item_pdf(out_pdf, summary)
    assert out_pdf.is_file()
    assert out_pdf.stat().st_size > 1000


def test_real_item_preview_v10_generation_smoke() -> None:
    item_root = Path(r"C:/Work/BARAKUDA_FULL/runs/ot/2026-04-08_164655/items/Gly20_brown_rep01")
    run_dir = item_root / "analysis"
    out_pdf = run_dir / "results" / "2026-04-09-OT-Gly20_brown_rep01-report-polished-preview-v10.pdf"
    if not run_dir.is_dir():
        pytest.skip("Real item folder is not available in this environment.")
    summary = build_ot_item_summary(
        run_dir=run_dir,
        base_name="Gly20_brown_rep01",
        item_id="Gly20_brown_rep01",
        source_input_path=str(item_root / "raw" / "Gly20_brown_rep01.raw"),
        status="success",
    )
    summary["analysis_dir"] = str(run_dir)
    ot_report.export_ot_item_pdf(out_pdf, summary)
    assert out_pdf.is_file()
    assert out_pdf.stat().st_size > 1000


def test_theory_page_lorentz_equation_replaced_with_stable_form() -> None:
    path = Path(ot_report.__file__)
    src = path.read_text(encoding="utf-8")
    assert r"S_{xx}(f)=\frac{A}{1+(f/f" in src
    assert r'$P(f) = \frac{A}{f_{\mathrm{c}}^{2} + f^{2}} + B$' not in src


def test_export_brownian_pdf_smoke(tmp_path: Path) -> None:
    run_dir = tmp_path / "analysis"
    base = "smoke"
    audit = run_dir / "audit"
    audit.mkdir(parents=True)
    _write(
        audit / "run.json",
        {"config": {"tracking": {"timing": {}}, "postprocess": {"physics_mode": "BROWNIAN"}, "calibration": {}}},
    )
    _write(audit / f"{base}_postprocess.json", {"summary": {"qc": {"lost_fraction": 0.0}}})
    _write(
        audit / f"{base}_calibration.json",
        {
            "kappa": {"kappa_x_pn_per_um": 1.0, "kappa_y_pn_per_um": 1.0},
            "viscosity": {"eta_mean_pa_s": 0.001},
            "diffusion": {"D_m2_s": 1e-12},
            "diagnostics": {"fc_x_hz": 10.0, "fc_y_hz": 10.0},
        },
    )
    summary = build_ot_item_summary(
        run_dir=run_dir, base_name=base, item_id=base, source_input_path="x.raw", status="success"
    )
    out = tmp_path / "out.pdf"
    ot_report.export_ot_item_pdf(out, summary)
    assert out.is_file() and out.stat().st_size > 1000


def test_scale_resolution_explicit_and_fallback() -> None:
    ok = resolve_scale_for_run(explicit_scale_um_per_px=0.06042, scale_source="dataset_sidecar")
    assert ok.um_per_px == pytest.approx(0.06042)
    assert ok.scale_source == "dataset_sidecar"
    assert ok.scale_warning is None

    missing = resolve_scale_for_run(explicit_scale_um_per_px=None, scale_source="none")
    assert missing.um_per_px is None
    assert missing.scale_warning is not None


def test_export_collision_protection(tmp_path: Path) -> None:
    exports_dir = tmp_path / "exports"
    first, coll1 = build_collision_safe_export_path(exports_dir, "results/summary.csv")
    first.parent.mkdir(parents=True, exist_ok=True)
    first.write_text("a", encoding="utf-8")
    second, coll2 = build_collision_safe_export_path(exports_dir, "results/summary.csv")
    assert coll1 is False
    assert coll2 is True
    assert second != first
    assert second.name.startswith("summary__dup")
